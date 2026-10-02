"""Document ingestion pipeline for the AI Library.

    Upload → Extract (pypdf / python-docx / python-pptx / csv / text) → Normalize → Chunk → Embed → pgvector
           → AI summary + classification
"""

import csv
import io
import re
import unicodedata
import uuid
from typing import Any

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.logging import get_logger
from app.models import Document, DocumentChunk, Organization
from app.schemas.ai import DocumentAnalysis
from app.services.ai import prompts
from app.services.ai.service import AIContext, ai_service
from app.services.embeddings import get_embedder
from app.services.storage import get_storage

log = get_logger(__name__)

CHUNK_CHARS = 1100
CHUNK_OVERLAP = 150


class ExtractionError(Exception):
    pass


def extract_text(data: bytes, ext: str) -> tuple[str, dict[str, Any]]:
    ext = ext.lower()
    meta: dict[str, Any] = {}
    if ext == "pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
        except Exception as e:
            raise ExtractionError(f"Unreadable PDF: {e}") from e
        pages = [(p.extract_text() or "") for p in reader.pages]
        meta["pages"] = len(pages)
        text = "\n\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(pages) if t.strip())
        if not text.strip():
            raise ExtractionError("PDF contains no extractable text (scanned images need OCR)")
        return text, meta
    if ext == "docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        meta["paragraphs"] = len(d.paragraphs)
        return "\n".join(parts), meta
    if ext == "pptx":
        from pptx import Presentation as Pptx

        prs = Pptx(io.BytesIO(data))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        texts.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text
                if notes.strip():
                    texts.append(f"Notes: {notes}")
            slides.append(f"[Slide {i}]\n" + "\n".join(t for t in texts if t.strip()))
        meta["slides"] = len(slides)
        return "\n\n".join(slides), meta
    if ext == "csv":
        text = data.decode("utf-8", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        meta["rows"] = len(rows)
        if not rows:
            return "", meta
        header = rows[0]
        lines = [", ".join(f"{h}: {v}" for h, v in zip(header, r, strict=False)) for r in rows[1:]]
        return "Columns: " + ", ".join(header) + "\n" + "\n".join(lines), meta
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="replace"), meta
    raise ExtractionError(f"No parser for .{ext}")


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"-\n(?=[a-z])", "", text)  # de-hyphenate line breaks from PDFs
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Paragraph-aware chunking with character overlap; long paragraphs are split on sentence boundaries."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    units: list[str] = []
    for p in paras:
        if len(p) <= size:
            units.append(p)
            continue
        sents = re.split(r"(?<=[.!?؟])\s+", p)
        buf = ""
        for s in sents:
            if len(buf) + len(s) + 1 > size and buf:
                units.append(buf)
                buf = ""
            while len(s) > size:
                units.append(s[:size])
                s = s[size:]
            buf = f"{buf} {s}".strip()
        if buf:
            units.append(buf)
    chunks: list[str] = []
    cur = ""
    for u in units:
        if len(cur) + len(u) + 2 > size and cur:
            chunks.append(cur)
            cur = cur[-overlap:] + "\n" + u if overlap else u
        else:
            cur = f"{cur}\n\n{u}" if cur else u
    if cur:
        chunks.append(cur)
    return chunks


async def ingest_document(session: AsyncSession, doc: Document, data: bytes | None = None) -> Document:
    """Runs the full pipeline for one document. Safe to re-run (replaces chunks)."""
    doc.status = "processing"
    await session.commit()
    try:
        if data is None:
            if not doc.storage_key:
                if not doc.extracted_text:
                    raise ExtractionError("Document has no stored file")
                raw_text, meta = doc.extracted_text, {}
            else:
                data = await get_storage().get(doc.storage_key)
        if data is not None:
            ext = doc.filename.rsplit(".", 1)[-1]
            raw_text, meta = extract_text(data, ext)
        text = normalize(raw_text)
        if len(text) < 20:
            raise ExtractionError("Document contains too little text to index")
        pieces = chunk_text(text)
        vectors = await get_embedder().embed(pieces)

        await session.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))
        for i, (content, vec) in enumerate(zip(pieces, vectors, strict=True)):
            session.add(
                DocumentChunk(
                    document_id=doc.id,
                    organization_id=doc.organization_id,
                    chunk_index=i,
                    content=content,
                    token_count=max(1, len(content) // 4),
                    embedding=vec,
                    meta={"title": doc.title, "category": doc.category},
                )
            )
        doc.extracted_text = text
        doc.meta = {**(doc.meta or {}), **meta, "chars": len(text), "chunks": len(pieces), "embedding": get_embedder().name}
        org = await session.get(Organization, doc.organization_id)
        org_settings = org.settings if org else None
        await session.commit()  # never hold a write transaction across an LLM call

        try:
            instructions, payload = prompts.document_analysis_prompt(filename=doc.filename, text=text)
            analysis = await ai_service.generate(
                task="document.analyze",
                schema=DocumentAnalysis,
                instructions=instructions,
                payload=payload,
                ctx=AIContext(organization_id=doc.organization_id, user_id=doc.uploaded_by, org_settings=org_settings),
                context={"text": text[:20000], "filename": doc.filename},
                temperature=0.1,
                max_tokens=1200,
            )
            doc.summary = analysis.summary
            doc.meta = {**doc.meta, "key_points": analysis.key_points, "suggested_category": analysis.category}
            if doc.category in ("", "auto", None):
                doc.category = analysis.category
        except Exception as e:  # summary is a nice-to-have; retrieval still works
            log.warning("document summary failed", extra={"data": {"document_id": str(doc.id), "error": str(e)[:200]}})
        doc.status = "ready"
        await session.commit()
    except Exception as e:
        await session.rollback()
        doc = await session.get(Document, doc.id) or doc
        doc.status = "failed"
        doc.meta = {**(doc.meta or {}), "error": str(e)[:500]}
        await session.commit()
        log.warning("document ingestion failed", extra={"data": {"document_id": str(doc.id), "error": str(e)[:300]}})
        raise
    return doc


async def retrieve(
    session: AsyncSession, org_id: uuid.UUID, query: str, k: int = 6, categories: list[str] | None = None
) -> list[dict[str, Any]]:
    """Top-k chunks for `query`, always scoped to one organization."""
    from app.core.config import get_settings
    from app.services.embeddings import cosine

    if not query.strip():
        return []
    [qvec] = await get_embedder().embed([query])
    settings = get_settings()
    base = (
        select(DocumentChunk, Document.title, Document.category, Document.id)
        .join(Document, Document.id == DocumentChunk.document_id)
        .where(DocumentChunk.organization_id == org_id, Document.status == "ready")
    )
    if categories:
        base = base.where(Document.category.in_(categories))

    if not settings.is_sqlite and settings.pgvector_enabled:
        from pgvector.sqlalchemy import Vector
        from sqlalchemy import Float, bindparam

        dist = DocumentChunk.embedding.op("<=>", return_type=Float)(
            bindparam("qvec", qvec, type_=Vector(settings.embedding_dim))
        )
        rows = (await session.execute(base.add_columns(dist.label("distance")).order_by(dist).limit(k))).all()
        return [
            {"content": c.content, "title": t, "category": cat, "document_id": str(did), "score": round(1 - float(d), 4)}
            for c, t, cat, did, d in rows
        ]

    rows = (await session.execute(base)).all()
    scored = [
        (cosine(qvec, c.embedding or []), c, t, cat, did) for c, t, cat, did in rows if c.embedding
    ]
    scored.sort(key=lambda x: x[0], reverse=True)
    return [
        {"content": c.content, "title": t, "category": cat, "document_id": str(did), "score": round(s, 4)}
        for s, c, t, cat, did in scored[:k]
    ]
