"""Editable PowerPoint export with python-pptx.

Every element becomes a native object — text boxes, shapes, connectors, PowerPoint charts (with
embedded data you can edit in PowerPoint), and native tables. Speaker notes are included and slide
transitions are mapped (fade → fade, dissolve → dissolve, slide → push). Element entrance animations
are not mapped (PowerPoint's timing XML is out of scope); the web presenter renders them.
"""

import io
from typing import Any

from lxml import etree
from pptx import Presentation as PptxPresentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.services.export.layout import H, W, Box, Canvas, Colors, SlideRenderer

CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_bar": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "donut": XL_CHART_TYPE.DOUGHNUT,
}
TRANSITIONS = {"fade": "<p:fade/>", "dissolve": "<p:dissolve/>", "slide": '<p:push dir="l"/>'}


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#")[:6].upper())


def _e(v: float) -> Emu:
    return Pt(v)


class PptxCanvas(Canvas):
    def __init__(self, slide: Any, fonts: dict[str, str]) -> None:
        self.slide = slide
        self.shapes = slide.shapes
        self.fonts = fonts

    def _font_name(self, font: str) -> str:
        if font == "heading":
            return self.fonts.get("heading", "Georgia")
        if font == "mono":
            return self.fonts.get("mono", "Consolas")
        return self.fonts.get("body", "Calibri")

    def rect(self, b: Box, fill: str, radius: float = 0, line: str | None = None) -> None:
        shp = self.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, _e(b.x), _e(b.y), _e(b.w), _e(b.h))
        if radius:
            shp.adjustments[0] = min(0.5, radius / max(1.0, min(b.w, b.h)))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if line:
            shp.line.color.rgb = _rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False

    def ellipse(self, b: Box, fill: str) -> None:
        shp = self.shapes.add_shape(MSO_SHAPE.OVAL, _e(b.x), _e(b.y), _e(b.w), _e(b.h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        shp.line.fill.background()
        shp.shadow.inherit = False

    def line(self, x1: float, y1: float, x2: float, y2: float, color: str, width: float = 1, arrow: bool = False) -> None:
        conn = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _e(x1), _e(y1), _e(x2), _e(y2))
        conn.line.color.rgb = _rgb(color)
        conn.line.width = Pt(width)
        if arrow:
            ln = conn.line._get_or_add_ln()
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")

    def text(self, b: Box, text: str, *, size: float, color: str, font: str = "body", bold: bool = False,
             align: str = "left", valign: str = "top", italic: bool = False) -> None:
        tb = self.shapes.add_textbox(_e(b.x), _e(b.y), _e(b.w), _e(max(b.h, size * 1.3)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(valign, MSO_ANCHOR.TOP)
        p = tf.paragraphs[0]
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
        run = p.add_run()
        run.text = text
        f = run.font
        f.size = Pt(size)
        f.bold = bold or font == "body_bold"
        f.italic = italic
        f.name = self._font_name("body" if font == "body_bold" else font)
        f.color.rgb = _rgb(color)

    def bullets(self, b: Box, items: list[str], *, size: float, color: str, bullet_color: str, ordered: bool = False) -> None:
        tb = self.shapes.add_textbox(_e(b.x), _e(b.y), _e(b.w), _e(b.h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(size * 0.6)
            marker = p.add_run()
            marker.text = f"{i + 1:02d}   " if ordered else "●  "
            marker.font.size = Pt(size * (0.75 if ordered else 0.6))
            marker.font.color.rgb = _rgb(bullet_color)
            marker.font.name = self._font_name("mono" if ordered else "body")
            r = p.add_run()
            r.text = item
            r.font.size = Pt(size)
            r.font.color.rgb = _rgb(color)
            r.font.name = self._font_name("body")

    def chart(self, b: Box, el: dict[str, Any], c: Colors) -> None:
        data = CategoryChartData()
        data.categories = el.get("labels", [])
        series = el.get("series", [])
        for s in series:
            data.add_series(s.get("name", "Series"), [float(v) for v in s.get("values", [])])
        ctype = CHART_TYPES.get(el.get("chart_type", "bar"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        gf = self.shapes.add_chart(ctype, _e(b.x), _e(b.y), _e(b.w), _e(b.h), data)
        chart = gf.chart
        chart.font.size = Pt(10)
        chart.font.color.rgb = _rgb(c.muted)
        chart.has_legend = len(series) > 1 or ctype in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        palette = [c.accent, c.line, c.primary, c.secondary]
        if ctype in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
            for i, point in enumerate(chart.plots[0].series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb([c.accent, c.primary, c.secondary, c.line, c.muted][i % 5])
            return
        hi = el.get("highlight_index")
        for si, s in enumerate(chart.plots[0].series):
            col = palette[si % len(palette)] if len(series) > 1 or hi is None else c.line
            if ctype in (XL_CHART_TYPE.LINE_MARKERS,):
                s.format.line.color.rgb = _rgb(c.accent if si == 0 else palette[si % len(palette)])
                s.format.line.width = Pt(2.5)
                s.smooth = False
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = _rgb(c.accent if (len(series) == 1 and hi is None) else col)
                if hi is not None and len(series) == 1 and 0 <= hi < len(el.get("labels", [])):
                    pt = s.points[hi]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = _rgb(c.accent)
        try:
            chart.value_axis.has_major_gridlines = False
            chart.value_axis.format.line.fill.background()
            chart.category_axis.format.line.color.rgb = _rgb(c.line)
        except (ValueError, AttributeError):
            pass

    def table(self, b: Box, columns: list[str], rows: list[list[str]], c: Colors) -> None:
        n_rows = len(rows) + 1
        h = min(b.h, 30 * n_rows)
        gf = self.shapes.add_table(n_rows, len(columns), _e(b.x), _e(b.y), _e(b.w), _e(h))
        tbl = gf.table
        for ci, name in enumerate(columns):
            cell = tbl.cell(0, ci)
            cell.text = name
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(c.fg if c.bg.upper() in ("#FFFFFF", "#F6F4EF") else c.panel)
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(11)
            para.runs[0].font.bold = True
            para.runs[0].font.color.rgb = _rgb(c.bg if c.bg.upper() in ("#FFFFFF", "#F6F4EF") else c.fg)
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(c.panel if ri % 2 == 0 else c.bg)
                run = cell.text_frame.paragraphs[0].runs[0]
                run.font.size = Pt(11)
                run.font.color.rgb = _rgb(c.fg)

    def image(self, b: Box, data: bytes | None, alt: str, c: Colors) -> None:
        if data:
            try:
                pic = self.shapes.add_picture(io.BytesIO(data), _e(b.x), _e(b.y), _e(b.w), _e(b.h))
                pic.name = alt[:60] or "Image"
                return
            except Exception:
                pass
        self.rect(b, c.panel, radius=8)
        self.ellipse(Box(b.x + b.w * 0.55, b.y + b.h * 0.5, b.w * 0.55, b.w * 0.55), c.accent)
        if alt:
            self.text(Box(b.x + 16, b.y + b.h - 30, b.w - 32, 18), alt, size=10, color=c.muted, font="mono")


def _set_transition(slide: Any, kind: str, duration_ms: int) -> None:
    xml = TRANSITIONS.get(kind)
    if not xml:
        return
    spd = "fast" if duration_ms < 400 else "slow" if duration_ms > 800 else "med"
    el = etree.fromstring(
        f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="{spd}">{xml}</p:transition>'
    )
    sld = slide._element
    # <p:transition> must come after <p:clrMapOvr> and before <p:timing>
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is not None:
        anchor.addnext(el)
    else:
        sld.find(qn("p:cSld")).addnext(el)


def build_pptx(deck: dict[str, Any], images: dict[str, bytes] | None = None, include_notes: bool = True) -> bytes:
    prs = PptxPresentation()
    prs.slide_width = Pt(W)
    prs.slide_height = Pt(H)
    blank = prs.slide_layouts[6]
    theme = deck.get("theme") or {}
    fonts = theme.get("fonts") or {}
    slides = deck.get("slides") or []
    prs.core_properties.title = deck.get("title", "Presentation")
    prs.core_properties.subject = deck.get("subtitle", "")
    for i, s in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        SlideRenderer(PptxCanvas(slide, fonts), theme, images).render(s, i, len(slides))
        if include_notes and s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        tr = s.get("transition") or {}
        _set_transition(slide, tr.get("type", "fade"), int(tr.get("duration_ms", 500)))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
