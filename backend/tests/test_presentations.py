"""Flow 1: create presentation → generate → edit → export."""

import io
import uuid

from pptx import Presentation as Pptx

PROMPT = "Create a 12-slide investor presentation about our fintech platform."


async def _generate(client, h, **kw):
    r = await client.post("/presentations/generate", json={"prompt": PROMPT, "deck_type": "investor", "audience": "investors", "slide_count": 12, **kw}, headers=h)
    assert r.status_code == 202, r.text
    body = r.json()
    assert body["job"]["status"] == "completed", body["job"]
    return (await client.get(f"/presentations/{body['presentation_id']}", headers=h)).json()


async def test_generate_edit_export(client, ws):
    h = ws["headers"]
    p = await _generate(client, h)
    deck = p["deck"]
    assert p["status"] == "draft"
    assert len(deck["slides"]) == 12
    assert deck["slides"][0]["layout"] == "hero" and deck["slides"][-1]["layout"] == "closing"
    assert deck["theme"]["palette"]["primary"] == "#2F5BD3"  # brand applied server-side

    # AI Designer edit on the current slide
    slide = deck["slides"][1]
    r = await client.post(f"/presentations/{p['id']}/ai-edit", json={"slide": slide, "command": "Make this more visual"}, headers=h)
    assert r.status_code == 200, r.text
    edited = r.json()["slide"]
    assert edited["id"] == slide["id"]
    assert edited["layout"] == "grid"

    # save edited deck (manual + AI edits) and rename
    deck["slides"][1] = edited
    deck["slides"].append({"layout": "quote", "title": "Customers love it", "elements": [{"type": "quote", "text": "Best tool we use.", "attribution": "CFO"}]})
    r = await client.patch(f"/presentations/{p['id']}", json={"deck": deck, "title": "Fintech Series B"}, headers=h)
    assert r.status_code == 200
    saved = r.json()
    assert saved["title"] == "Fintech Series B" and saved["slide_count"] == 13

    # editable PPTX
    job = (await client.post(f"/presentations/{p['id']}/export", json={"format": "pptx"}, headers=h)).json()
    assert job["status"] == "completed", job
    data = (await client.get(f"/jobs/{job['id']}/download", headers=h)).content
    prs = Pptx(io.BytesIO(data))
    assert len(prs.slides) == 13
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert any("Fintech" in t for t in texts)  # real text, not images
    assert any(sh.has_chart for s in prs.slides for sh in s.shapes)  # native, editable charts
    assert prs.slides[0].has_notes_slide

    # PDF
    job = (await client.post(f"/presentations/{p['id']}/export", json={"format": "pdf"}, headers=h)).json()
    pdf = (await client.get(f"/jobs/{job['id']}/download", headers=h)).content
    assert pdf.startswith(b"%PDF")


async def test_invalid_deck_rejected(client, ws):
    p = (await client.post("/presentations", json={"title": "x"}, headers=ws["headers"])).json()
    bad = {"title": "x", "slides": [{"layout": "not-a-layout", "title": "x"}]}
    r = await client.patch(f"/presentations/{p['id']}", json={"deck": bad}, headers=ws["headers"])
    assert r.status_code == 422


async def test_crud_duplicate_favorite_style_folder(client, ws):
    h = ws["headers"]
    p = await _generate(client, h, slide_count=5, deck_type="board")
    dup = (await client.post(f"/presentations/{p['id']}/duplicate", headers=h)).json()
    assert dup["title"].endswith("(copy)") and dup["slide_count"] == 5
    await client.put(f"/presentations/{p['id']}/favorite", headers=h)
    favs = (await client.get("/presentations", params={"favorites": True}, headers=h)).json()
    assert [f["id"] for f in favs] == [p["id"]]
    styled = (await client.patch(f"/presentations/{p['id']}", json={"style": "editorial"}, headers=h)).json()
    assert styled["deck"]["theme"]["style"] == "editorial"
    folder = (await client.post("/folders", json={"name": "Board"}, headers=h)).json()
    await client.patch(f"/presentations/{p['id']}", json={"folder_id": folder["id"]}, headers=h)
    assert len((await client.get("/presentations", params={"folder_id": folder["id"]}, headers=h)).json()) == 1
    found = (await client.get("/presentations", params={"q": "copy"}, headers=h)).json()
    assert [f["id"] for f in found] == [dup["id"]]
    assert (await client.delete(f"/presentations/{dup['id']}", headers=h)).status_code == 204
    assert (await client.get(f"/presentations/{dup['id']}", headers=h)).status_code == 404


async def test_template_flow(client, ws):
    h = ws["headers"]
    p = await _generate(client, h, slide_count=6)
    t = (await client.post("/templates", json={"name": "Our investor deck", "category": "investor", "presentation_id": p["id"]}, headers=h)).json()
    assert t["slide_count"] == 6
    new = (await client.post("/presentations", json={"title": "From template", "template_id": t["id"]}, headers=h)).json()
    assert new["slide_count"] == 6 and new["source"] == "template"
    assert (await client.get(f"/templates/{t['id']}", headers=h)).json()["usage_count"] == 1


async def test_saved_slide_ai_edit(client, ws):
    from sqlalchemy import select

    from app.db.session import SessionLocal
    from app.models import PresentationSlide

    p = await _generate(client, ws["headers"], slide_count=5)
    async with SessionLocal() as s:
        row = (await s.execute(select(PresentationSlide).where(PresentationSlide.presentation_id == uuid.UUID(p["id"]), PresentationSlide.position == 1))).scalar_one()
    r = await client.post(f"/slides/{row.id}/ai-edit", json={"command": "Add a chart"}, headers=ws["headers"])
    assert r.status_code == 200
    again = (await client.get(f"/presentations/{p['id']}", headers=ws["headers"])).json()
    assert any(e["type"] == "chart" for e in again["deck"]["slides"][1]["elements"])
