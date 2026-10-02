"""Flow 5: upload company document → process → use in AI generation. Plus sharing and upload validation."""

import io
from datetime import timedelta

import docx

from tests.conftest import make_workspace


def _docx_bytes(text: str) -> bytes:
    d = docx.Document()
    for para in text.split("\n"):
        d.add_paragraph(para)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def test_docx_ingestion_retrieval_and_generation(client, ws):
    h = ws["headers"]
    body = _docx_bytes("Services overview\nWe provide quantum logistics optimisation for airlines.\nResults: 17% fuel savings for Skyway Air in 2025.")
    r = await client.post("/documents", files={"file": ("services.docx", io.BytesIO(body), "application/octet-stream")}, data={"category": "services"}, headers=h)
    assert r.status_code == 202, r.text
    d = (await client.get(f"/documents/{r.json()['id']}", headers=h)).json()
    assert d["status"] == "ready" and d["chunk_count"] >= 1
    assert "quantum logistics" in d["extracted_text"]
    assert d["summary"]
    hits = (await client.post("/documents/search", json={"query": "quantum logistics airlines fuel"}, headers=h)).json()
    assert hits and hits[0]["document_id"] == d["id"]
    gen = (await client.post("/presentations/generate", json={"prompt": "Sales deck about quantum logistics for airlines", "deck_type": "sales", "slide_count": 6}, headers=h)).json()
    assert "services" in gen["job"]["result"]["sources"]


async def test_upload_validation(client, ws):
    h = ws["headers"]
    fake_pdf = await client.post("/documents", files={"file": ("x.pdf", io.BytesIO(b"not a pdf at all"), "application/pdf")}, headers=h)
    assert fake_pdf.status_code == 415
    exe = await client.post("/assets", files={"file": ("run.exe", io.BytesIO(b"MZ..."), "application/octet-stream")}, headers=h)
    assert exe.status_code == 415
    svg = await client.post("/assets", files={"file": ("x.svg", io.BytesIO(b"<svg><script>alert(1)</script></svg>"), "image/svg+xml")}, headers=h)
    assert svg.status_code == 415
    empty = await client.post("/documents", files={"file": ("x.txt", io.BytesIO(b""), "text/plain")}, headers=h)
    assert empty.status_code == 400


async def test_share_links_password_expiry_internal(client):
    a = await make_workspace(client, "Sharer")
    b = await make_workspace(client, "Outsider")
    p = (await client.post("/presentations", json={"title": "Shared deck"}, headers=a["headers"])).json()

    ext = (await client.post(f"/presentations/{p['id']}/share", json={"audience": "external", "password": "open-sesame", "allow_download": False}, headers=a["headers"])).json()
    assert (await client.get(f"/public/shares/{ext['token']}")).status_code == 401
    assert (await client.get(f"/public/shares/{ext['token']}", headers={"X-Share-Password": "nope"})).status_code == 401
    ok = await client.get(f"/public/shares/{ext['token']}", headers={"X-Share-Password": "open-sesame"})
    assert ok.status_code == 200 and ok.json()["deck"]["title"] == "Shared deck"
    assert (await client.get(f"/public/shares/{ext['token']}/download", headers={"X-Share-Password": "open-sesame"})).status_code == 403

    internal = (await client.post(f"/presentations/{p['id']}/share", json={"audience": "internal", "allow_download": True}, headers=a["headers"])).json()
    assert (await client.get(f"/public/shares/{internal['token']}")).status_code == 401
    assert (await client.get(f"/public/shares/{internal['token']}", headers=b["headers"])).status_code == 403
    assert (await client.get(f"/public/shares/{internal['token']}", headers=a["headers"])).status_code == 200
    dl = await client.get(f"/public/shares/{internal['token']}/download", params={"format": "pptx"}, headers=a["headers"])
    assert dl.status_code == 200 and dl.content[:2] == b"PK"

    # expiry + revoke
    from app.db.base import utcnow
    from app.db.session import SessionLocal
    from app.models import ShareLink

    async with SessionLocal() as s:
        link = await s.get(ShareLink, __import__("uuid").UUID(ext["id"]))
        link.expires_at = utcnow() - timedelta(minutes=1)
        await s.commit()
    assert (await client.get(f"/public/shares/{ext['token']}", headers={"X-Share-Password": "open-sesame"})).status_code == 410
    assert (await client.delete(f"/public/shares/{internal['id']}", headers=a["headers"])).status_code == 204
    assert (await client.get(f"/public/shares/{internal['token']}", headers=a["headers"])).status_code == 404
